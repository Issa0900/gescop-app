# -*- coding: utf-8 -*-
"""
Vérité « tous modules » des fichiers du dossier DEMO, calculée ICI directement
depuis les fichiers, sans le moteur de GESCOP. Complète verite_demo.ts (qui ne
contrôlait guère que le CA et le nombre de lignes) : pour chaque module présent
dans un classeur (ventes, clients, produits, stocks, employés, paie, dépenses,
trésorerie, fournisseurs, achats, immobilisations, marketing, interactions,
concurrents, veille, objectifs, événements), le nombre de lignes et les
chiffres que la page du module doit afficher.

Règles de calcul (celles de l'app, écrites noir sur blanc) :
  - paie : coût employeur total de chaque fiche (payroll_total) ;
  - dépenses : montant hors taxes (les taxes payées sont récupérables) ;
  - achats : coût hors taxes (sous-total) ;
  - stocks : valeur au coût (colonne de valeur, sinon quantité × coût unitaire) ;
  - immobilisations : valeur nette comptable ;
  - trésorerie : solde de clôture de la dernière période ;
  - clients actifs : statut « Actif » ;
  - listes (canal, sentiment, famille…) : valeur de l'app après traduction.

Usage : python tests/banc/verite_demo_modules.py  (écrit verite_demo_modules.json)
"""
import json, os
from collections import Counter
from openpyxl import load_workbook

ICI = os.path.dirname(os.path.abspath(__file__))
DEMO = os.path.abspath(os.path.join(ICI, "..", "..", "..", "DEMO"))

def arr(x): return round(x + 1e-9, 2)
def num(v):
    if v is None or v == "": return 0.0
    return float(v)

def feuille(wb, nom, ignorer_prefixe="TOTAL"):
    ws = wb[nom]
    rows = list(ws.iter_rows(values_only=True))
    idx_tete = 0
    for idx, r in enumerate(rows):
        non_vides = [c for c in r if c is not None and str(c).strip() != ""]
        if len(non_vides) >= 3:
            idx_tete = idx
            break
    tete = [str(h).strip() if h is not None else "" for h in rows[idx_tete]]
    out = []
    for r in rows[idx_tete + 1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r): continue
        prem = str(r[0] or "").strip().upper()
        if ignorer_prefixe and prem.startswith(ignorer_prefixe): continue
        out.append({tete[i]: r[i] for i in range(len(tete)) if i < len(r) and tete[i]})
    return out

def derniere_photo(rows, produit, entrepot, date):
    """Un stock est une photo : pour un meme produit dans un meme entrepot, seule
    la ligne la plus recente compte (additionner les dates double le stock)."""
    last = {}
    for r in rows:
        k = (r[produit], r[entrepot])
        if k not in last or str(r[date]) >= str(last[k][date]): last[k] = r
    return list(last.values())

def kpi(id_, attendu, note=""): return {"type": "kpi", "id": id_, "attendu": attendu, "note": note}
def lignes(entite, n): return {"type": "lignes", "entite": entite, "attendu": n}
def somme(entite, champ, v, note=""): return {"type": "somme", "entite": entite, "champ": champ, "attendu": arr(v), "note": note}
def distincts(entite, champ, n): return {"type": "distincts", "entite": entite, "champ": champ, "attendu": n}
def compte(entite, champ, valeur, n): return {"type": "compte", "entite": entite, "champ": champ, "valeur": valeur, "attendu": n}
def solde(v): return {"type": "solde", "attendu": arr(v)}

VERITE = []

# ------------------------------------------------------------ Xplorer 3 mois (18 feuilles, une par module)
def xplorer_3mois():
    f = "GESCOP_Donnees_Test_Xplorer_3Mois.xlsx"
    wb = load_workbook(os.path.join(DEMO, f), read_only=True, data_only=True)
    F = lambda n: feuille(wb, n)
    pay, exp, pur, prod = F("Payroll"), F("Expenses"), F("Purchase"), F("Products")
    cus, inter, emp, cf = F("Customers"), F("Interaction"), F("Employees"), F("Cashflow")
    camp, cd, comp, sig, goal, ev, sup, inv = F("Campaigns"), F("CampaignDaily"), F("Competitors"), F("ExternalSignal"), F("Goal"), F("Event"), F("Suppliers"), F("Inventory")
    dernier = max(cf, key=lambda r: str(r["date"]))
    senti = Counter(str(r["sentiment"]).lower() for r in inter)
    VERITE.append({"fichier": f, "controles": [
        lignes("Payroll", len(pay)), kpi("payroll_total", arr(sum(num(r["total_cost"]) for r in pay))),
        lignes("Expense", len(exp)), somme("Expense", "amount", sum(num(r["amount"]) for r in exp)),
        lignes("Purchase", len(pur)), somme("Purchase", "total_cost", sum(num(r["total_cost"]) for r in pur)),
        lignes("Product", len(prod)), somme("Product", "selling_price", sum(num(r["selling_price"]) for r in prod)),
        lignes("Supplier", len(sup)),
        lignes("Inventory", len(inv)),
        lignes("Customer", len(cus)), kpi("active_customers", sum(1 for r in cus if str(r["status"]).lower() == "actif")),
        lignes("Interaction", len(inter)), compte("Interaction", "sentiment", "positif", senti.get("positif", 0)),
        compte("Interaction", "sentiment", "negatif", senti.get("négatif", 0) + senti.get("negatif", 0)),
        lignes("Employee", len(emp)),
        lignes("Cashflow", len(cf)), solde(num(dernier["closing_cash"])),
        lignes("Campaign", len(camp)), somme("Campaign", "spend", sum(num(r["spend"]) for r in camp)),
        lignes("CampaignDaily", len(cd)), somme("CampaignDaily", "spend", sum(num(r["spend"]) for r in cd)),
        lignes("Competitor", len(comp)), compte("Competitor", "market_position", "leader", sum(1 for r in comp if str(r["market_position"]).lower() == "leader")),
        lignes("ExternalSignal", len(sig)),
        lignes("Goal", len(goal)), somme("Goal", "target", sum(num(r["target"]) for r in goal)),
        lignes("Event", len(ev)),
    ]})

# ------------------------------------------------------------ Simulation Québec 3 ans (14 feuilles, suffixes _cad)
def simulation_3ans():
    f = "Simulation_Entreprise_Quebec_3Ans_Complet.xlsx"
    wb = load_workbook(os.path.join(DEMO, f), read_only=True, data_only=True)
    F = lambda n: feuille(wb, n)
    pay, exp, pur, ast, inv = F("Payroll"), F("Expense"), F("Purchase"), F("Asset"), F("Inventory")
    cus, emp, cf, cd, camp, prod, sup = F("Customer"), F("Employee"), F("Cashflow"), F("CampaignDaily"), F("Campaign"), F("Product"), F("Supplier")
    dernier = max(cf, key=lambda r: str(r["period_month"]))
    VERITE.append({"fichier": f, "controles": [
        kpi("payroll_total", arr(sum(num(r["total_employer_cost_cad"]) for r in pay)), "coût employeur total"),
        somme("Expense", "amount", sum(num(r["amount_ht_cad"]) for r in exp), "hors taxes"),
        somme("Purchase", "total_cost", sum(num(r["subtotal_cad"]) for r in pur), "hors taxes"),
        kpi("net_book_value_total", arr(sum(num(r["net_book_value_cad"]) for r in ast))),
        somme("Asset", "initial_cost", sum(num(r["acquisition_cost_cad"]) for r in ast)),
        lignes("Inventory", len(inv)), kpi("inventory_value_total", arr(sum(num(r["total_inventory_value_cad"]) for r in inv))),
        lignes("Customer", len(cus)), kpi("active_customers", sum(1 for r in cus if str(r["status"]).lower() == "actif")),
        lignes("Employee", len(emp)),
        lignes("Product", len(prod)), lignes("Supplier", len(sup)), lignes("Campaign", len(camp)),
        somme("Campaign", "budget", sum(num(r["allocated_budget_cad"]) for r in camp)),
        kpi("marketing_spend", arr(sum(num(r["daily_cost_cad"]) for r in cd)), "coût journalier des campagnes"),
        solde(num(dernier["closing_bank_balance_cad"])),
    ]})

# ------------------------------------------------------------ GESCOP.xlsx (10 feuilles de 500 lignes, en-têtes français)
def gescop():
    f = "GESCOP.xlsx"
    wb = load_workbook(os.path.join(DEMO, f), read_only=True, data_only=True)
    F = lambda n: feuille(wb, n)
    emp, dep, inv, pur, mkt, tre, four, prod = F("Employes"), F("Depenses"), F("Inventaire_Stocks"), F("Achats_Fournisseurs"), F("Marketing"), F("Tresorerie"), F("Fournisseurs"), F("Produits")
    dernier = max(tre, key=lambda r: str(r["Date"]))
    esg = [num(r["Score ESG"]) for r in four if r.get("Score ESG") is not None]
    VERITE.append({"fichier": f, "controles": [
        kpi("total_expense", arr(sum(num(r["Montant"]) for r in dep)), "dépenses seules (aucune transaction)"),
        somme("Employee", "annual_salary", sum(num(r["Salaire annuel"]) for r in emp)),
        somme("Employee", "total_employer_cost", sum(num(r["Coût employeur total"]) for r in emp)),
        kpi("inventory_value_total", arr(sum(num(r["Valeur du stock"]) for r in derniere_photo(inv, "ID Produit", "ID Entrepôt", "Date"))),
            "valeur au coût, dernière photo par produit et entrepôt (un stock n'est pas un flux)"),
        somme("Purchase", "total_cost", sum(num(r["Coût total"]) for r in pur)),
        kpi("marketing_spend", arr(sum(num(r["Dépense"]) for r in mkt))),
        kpi("weighted_esg_score", arr(sum(esg) / len(esg))),
        somme("Product", "selling_price", sum(num(r["Prix de vente"]) for r in prod)),
        solde(num(dernier["Solde de clôture"])),
    ]})

# ------------------------------------------------------------ Nordik Plein Air (7 feuilles, en-têtes avec titres)
def nordik():
    f = "Nordik_PleinAir_Donnees_Complet_2026.xlsx"
    wb = load_workbook(os.path.join(DEMO, f), read_only=True, data_only=True)
    F = lambda n: feuille(wb, n)
    inv, crm, mkt, emp, sup = F("Stocks & Inventaire"), F("Clients (CRM)"), F("Marketing"), F("Employés & RH"), F("Fournisseurs")
    val_cout = sum(num(r["Qté en Stock"]) * num(r["Coût Unitaire ($)"]) for r in inv)
    val_vente = sum(num(r["Qté en Stock"]) * num(r["Prix Vente ($)"]) for r in inv)
    pts = sum(num(r["Points Fidélité"]) for r in crm)
    budget = sum(num(r["Budget (CAD)"]) for r in mkt)
    clics = sum(num(r["Clics"]) for r in mkt)
    succursales = len(set(r["Succursale"] for r in emp if r.get("Succursale")))
    VERITE.append({"fichier": f, "controles": [
        kpi("inventory_value_total", arr(val_cout), "valeur stock au coût"),
        somme("Inventory", "inventory_value", val_cout, "valeur stock coût"),
        somme("Inventory", "sale_value", val_vente, "valeur stock vente"),
        somme("Customer", "loyalty_points", pts, "points fidélité"),
        somme("Campaign", "budget", budget, "budget marketing"),
        somme("Campaign", "clicks", clics, "clics marketing"),
        distincts("Employee", "branch", succursales),
        lignes("Supplier", len(sup)),
    ]})

# ------------------------------------------------------------ Entreprise Simulation 50 Ans Canada QC (7 feuilles)
def simulation_50ans():
    f = "Entreprise_Simulation_50Ans_Canada_QC.xlsx"
    wb = load_workbook(os.path.join(DEMO, f), read_only=True, data_only=True)
    F = lambda n: feuille(wb, n)
    ast, inv, emp, crm, four = F("Registre_Immobilisations_50Ans"), F("Stocks_MultiEntrepots"), F("Employes_RH"), F("Clients_CRM"), F("Fournisseurs")
    cost_ast = sum(num(r["Cout_Acquisition_Initial_CAD"]) for r in ast)
    amort_ast = sum(num(r["Amortissement_Cumule_CAD"]) for r in ast)
    vnc_ast = cost_ast - amort_ast
    val_inv = sum(num(r["Quantite_En_Stock"]) * num(r["Cout_Moyen_Pondere_CAD"]) for r in inv)
    qte_stock = sum(num(r["Quantite_En_Stock"]) for r in inv)
    sal_emp = sum(num(r["Salaire_Base_Annuel_CAD"]) for r in emp)
    taux = [
        num(r["Cotisation_RRQ_Patronale"]) + num(r["Cotisation_RQAP_Patronale"]) +
        num(r["Cotisation_CNESST"]) + num(r["Cotisation_FSS_QC"]) +
        num(r["Assurance_Collective_Part_Patronale"]) + num(r["Regime_REER_Collectif_Employeur"])
        for r in emp
    ]
    cout_emp = sal_emp + sum(num(r["Salaire_Base_Annuel_CAD"]) * t for r, t in zip(emp, taux))
    pts_crm = sum(num(r["Solde_Points_Fidelite"]) for r in crm)
    credit_crm = sum(num(r["Limite_Credit_CAD"]) for r in crm)
    esg = [num(r["Score_RSE_ESG"]) for r in four if r.get("Score_RSE_ESG") is not None]
    avg_esg = sum(esg) / len(esg) if esg else 0.0
    VERITE.append({"fichier": f, "controles": [
        kpi("net_book_value_total", arr(vnc_ast), "valeur nette comptable"),
        somme("Asset", "initial_cost", cost_ast),
        somme("Asset", "accumulated_depreciation", amort_ast),
        somme("Asset", "net_book_value", vnc_ast),
        kpi("inventory_value_total", arr(val_inv), "valeur stock au coût"),
        somme("Inventory", "closing_stock", qte_stock),
        somme("Employee", "annual_salary", sal_emp),
        somme("Employee", "total_employer_cost", arr(cout_emp)),
        somme("Customer", "loyalty_points", pts_crm),
        somme("Customer", "credit_limit", credit_crm),
        kpi("weighted_esg_score", arr(avg_esg)),
        lignes("Supplier", len(four)),
    ]})

for fn in [xplorer_3mois, simulation_3ans, gescop, nordik, simulation_50ans]: fn()
with open(os.path.join(ICI, "verite_demo_modules.json"), "w", encoding="utf-8") as fh:
    json.dump(VERITE, fh, ensure_ascii=False, indent=1)
print(sum(len(v["controles"]) for v in VERITE), "contrôles pour", len(VERITE), "fichiers")
